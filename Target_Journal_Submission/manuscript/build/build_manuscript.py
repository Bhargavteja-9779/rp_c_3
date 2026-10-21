#!/usr/bin/env python3
"""Build the JKSU-CIS manuscript (anonymised DOCX + PDF) from machine-generated results.

Every number in the text is produced by evaluating a {{ expression }} against
code/results/statistics.json, code/results/metrics.json and the CSV tables; a missing
value aborts the build (nothing is typed by hand). Tables are generated from the same
files. Citations ([@key]) are resolved by pandoc citeproc with the Springer Basic
(author-date) style from references.json.
"""
import json
import math
import re
import subprocess
import sys
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
SUB = HERE.parents[1]                       # Target_Journal_Submission/
CODE = SUB / "code"
OUT = SUB / "manuscript"
ARGS = [a for a in sys.argv[1:] if not a.startswith("--")]
JOURNAL = "eswa" if "--journal=eswa" in sys.argv else "jksu"
RES = Path(ARGS[0]) if ARGS else CODE   # root holding results/, tables/, figures/
import os
DRAFT = os.environ.get("DRAFT") == "1"

S = json.load(open(RES / "results" / "statistics.json"))
M = json.load(open(RES / "results" / "metrics.json"))
T = {p.stem: p for p in (RES / "tables").glob("*.csv")}
_lr = RES / "results" / "check_logreg_convergence.json"
if not _lr.exists():
    _lr = CODE / "results" / "check_logreg_convergence.json"
LR = {}
if _lr.exists():
    _d = json.load(open(_lr))
    _g = sorted(_d, key=lambda r: r["nll_long"] - r["nll_default"])[0]
    LR = dict(med_def=float(np.median([r["nll_default"] for r in _d])), med_long=float(np.median([r["nll_long"] for r in _d])),
              max_gain=_g["nll_default"] - _g["nll_long"], max_ds=_g["dataset"], n=len(_d))


# ------------------------------------------------------------------ formatting helpers
def f(x, nd=3):
    if x is None or (isinstance(x, float) and math.isnan(x)):
        raise ValueError("missing value in manuscript expression")
    return f"{x:.{nd}f}"


def f2(x): return f(x, 2)
def f1(x): return f(x, 1)
def pct(x, nd=1): return f"{100 * x:.{nd}f} %"


def pv(p):
    return "< 0.001" if p < 1e-3 else f"= {p:.3f}"


def pw(key, metric, comp):
    """pairwise record of LOOM vs competitor"""
    for r in S[key][metric]["pairwise"]:
        if r["competitor"] == comp:
            return r
    raise KeyError((key, metric, comp))


def med(key, metric, method):
    return S[key][metric]["medians"][method]


def rank(key, metric, method):
    return S[key][metric]["mean_ranks"][method]


def wdl(key, metric, comp):
    r = pw(key, metric, comp)
    return f"{r['ref_wins']}/{r['draws']}/{r['ref_losses']}"


def csv(name):
    return pd.read_csv(T[name])


def md_table(df, caption, label=None, colfmt=None):
    cols = list(df.columns)
    lines = ["| " + " | ".join(cols) + " |", "|" + "|".join(["---"] * len(cols)) + "|"]
    for _, r in df.iterrows():
        lines.append("| " + " | ".join(str(r[c]) for c in cols) + " |")
    return "\n".join(lines) + f"\n\n: {caption}\n"


def tab_compare(key, metrics, methods, caption, names=None, lower=None):
    """One row per method: median of each metric, mean rank (first metric), W/D/L and Holm p
    (first metric) of LOOM vs method."""
    lower = lower or {}
    m0 = metrics[0]
    rows = []
    for m in methods:
        if m not in S[key][m0]["medians"]:
            continue
        row = {"Method": (names or NAMES).get(m, m), "n": str(S[key][m0]["n_available"][m])}
        for mt in metrics:
            if mt in S[key] and m in S[key][mt]["medians"]:
                row[mt.upper() if mt != "shift_auroc" else "Shift AUROC"] = f(S[key][mt]["medians"][m])
        row["Mean rank"] = f2(S[key][m0]["mean_ranks"][m]) if m in S[key][m0]["mean_ranks"] else "–"
        if m == "LOOM":
            row["W/D/L"] = "–"; row["r_rb"] = "–"; row["p (Holm)"] = "–"
        else:
            r = pw(key, m0, m)
            row["W/D/L"] = f"{r['ref_wins']}/{r['draws']}/{r['ref_losses']}"
            row["r_rb"] = f2(r["rank_biserial"])
            row["p (Holm)"] = pv(r["p_holm"]).replace("= ", "")
        rows.append(row)
    df = pd.DataFrame(rows)
    return md_table(df, caption)


def cmp(key, metric, comp, lower=True, name=None):
    """Data-driven comparison phrase of LOOM vs `comp` (direction and significance from the stats)."""
    r = pw(key, metric, comp)
    nm = name or NAMES.get(comp, comp)
    better = (r["median_diff"] > 0) if lower else (r["median_diff"] < 0)
    stats = (f"W/D/L {r['ref_wins']}/{r['draws']}/{r['ref_losses']}; Δ = {r['median_diff']:+.3f} "
             f"[{r['ci_low']:+.3f}, {r['ci_high']:+.3f}]; r_rb = {r['rank_biserial']:.2f}; p {pv(r['p_holm'])}")
    if r["p_holm"] < 0.05:
        return f"{'better' if better else 'worse'} than {nm} ({stats})"
    return f"not significantly different from {nm} ({stats})"


def sigword(key, metric, comp, lower=True):
    r = pw(key, metric, comp)
    better = (r["median_diff"] > 0) if lower else (r["median_diff"] < 0)
    if r["p_holm"] >= 0.05:
        return "not significantly different"
    return "significantly better" if better else "significantly worse"


def dd(key, metric, comp):
    """Compact paired statistic: median difference (competitor - LOOM) and Holm-adjusted p."""
    r = pw(key, metric, comp)
    return f"Δ = {r['median_diff']:+.3f}, p {pv(r['p_holm'])}"


def ts(frac, cond, comp, field):
    for r in S["E7_train_size"]:
        if r["train_frac"] == frac and r["cond"] == cond and r["competitor"] == comp:
            return r[field]
    raise KeyError((frac, cond, comp))


def tsp(frac, cond, comp="PreVal"):
    """Train-size comparison phrase: medians, W/D/L, Holm p."""
    return (f"median NLL {ts(frac, cond, comp, 'median_loom'):.3f} vs {ts(frac, cond, comp, 'median_comp'):.3f}, "
            f"W/D/L {ts(frac, cond, comp, 'wdl')}, p {pv(ts(frac, cond, comp, 'p_holm'))}, n = {ts(frac, cond, comp, 'n')}")


def tsclaim():
    ok = all(ts(fr, "shifted", "PreVal", "p_holm") < 0.05 and
             ts(fr, "shifted", "PreVal", "median_loom") < ts(fr, "shifted", "PreVal", "median_comp") for fr in (0.25, 0.5))
    return ("The advantage under corruption is therefore retained with reduced training sets." if ok else
            "The advantage under corruption is therefore not consistently retained with reduced training sets.")


def both(w1, w2):
    return f"both {w1}" if w1 == w2 else f"{w1} and {w2}, respectively,"


def varword(key, metric, comp, lower=True):
    """Like sigword, but from the perspective of the competitor/variant (the sentence subject)."""
    w = sigword(key, metric, comp, lower)
    return {"significantly better": "significantly worse", "significantly worse": "significantly better"}.get(w, w)


def cp(method, shifted=False, field="cov_mean"):
    for r in S["E9_conformal"]:
        if r["method"] == method and bool(r["shifted"]) == shifted:
            return r[field]
    raise KeyError((method, shifted))


def eff(method, field="fit_median"):
    for r in S["E6_efficiency"]:
        if r["method"] == method:
            return r[field]
    raise KeyError(method)


def ds_total(name):
    d = pd.read_csv(CODE / "data" / "ucr_equal_length.csv").set_index("dataset")
    return int(d.loc[name, "n_train"] + d.loc[name, "n_test"])


def ns(method, field):
    return S["E10_natural_shift"]["means"][method][field]


def nsw(metric, comp):
    for r in S["E10_natural_shift"][f"pairwise_{metric}"]:
        if r["competitor"] == comp:
            return r
    raise KeyError((metric, comp))


def nslev():
    v = list(S["E10_natural_shift"]["lev_ratio"].values())
    return min(v), max(v)


def pwg(cond, metric, comp):
    for r in S[f"sens_grouped_{cond}"][metric]["pairwise"]:
        if r["competitor"] == comp:
            return r
    raise KeyError((cond, metric, comp))


def winrate(cond, sev):
    for r in S["E8_win_rate_by_condition"]:
        if r["cond"] == cond and int(r["sev"]) == sev:
            return r["loom_better"]
    raise KeyError((cond, sev))


def corr(key, field="rho"):
    return S["E8_correlations"][key][field]


def worst(kind, i, field="dataset"):
    return S["E8_worst_clean" if kind == "clean" else "E8_worst_shift"][i][field]


FAMS = ["noise", "wander", "missing", "spikes", "warp", "shift"]


def monotone_except():
    """Families whose LOOM win rate does NOT increase strictly with severity."""
    return [c for c in FAMS if not (winrate(c, 1) < winrate(c, 2) < winrate(c, 3))]


def trendclaim():
    bad = monotone_except()
    if not bad:
        return "grows with severity for every corruption family"
    names = {"noise": "Gaussian noise", "wander": "baseline wander", "missing": "missing segments",
             "spikes": "spikes", "warp": "time warping", "shift": "circular shifts"}
    return "grows with severity for every family except " + " and ".join(names[c] for c in bad)


def sev1range(exclude=("warp",)):
    v = [winrate(c, 1) for c in FAMS if c not in exclude]
    return f"{min(v)*100:.0f} %–{max(v)*100:.0f} %"


def worstlist(kind, k):
    return ", ".join(worst(kind, i) for i in range(k - 1)) + " and " + worst(kind, k - 1)


def perclass_range(kind, k):
    v = [worst(kind, i, "n_train") / worst(kind, i, "n_classes") for i in range(k)]
    return f"{min(v):.0f} to {max(v):.0f}"


def lev(cond):
    return S["leverage_ratio_median_by_condition"][f"h_{cond}"]


def ncomp(key, metric="nll"):
    return S[key][metric]["n_datasets"]


NAMES = {"LOOM": "LOOM (ours)", "PreVal": "PreVal", "Ridge-softmax": "Ridge softmax (T = 1)", "Ridge-onehot": "Ridge one-hot",
         "Holdout-TS": "Hold-out temperature scaling", "Ridge-Platt-CV5": "Ridge + Platt (5-fold)",
         "Ridge-Isotonic-CV5": "Ridge + isotonic (5-fold)", "LogReg-CV": "Logistic regression (CV)", "QUANT": "QUANT",
         "Arsenal": "Arsenal", "MiniRocket-transform": "MiniRocket transform of the training set (shared)", "PreVal-Ens5": "PreVal, 5-seed ensemble", "LOOM-Ens5": "LOOM, 5-seed ensemble",
         "PreVal-smooth": "No leverage (PreVal + smoothing)", "LOOM-nosmooth": "No target smoothing",
         "LOOM-free-b": "Free moderation b", "LOOM-finite-c1": "Finite moderation (c = 1)", "LOOM-kNN": "k-NN distance instead of leverage",
         "LOOM-insample": "In-sample instead of LOO pairs", "LOOM-holdout": "80/20 hold-out instead of LOO pairs",
         "LOOM-lambdaSSE": "λ by LOO squared error"}


def tab_conformal(caption):
    rows = []
    for m, lab in [("CP-LOOM-jackknife+", "LOOM, jackknife+, LAC"), ("CP-LOOM-jackknife", "LOOM, jackknife, LAC"),
                   ("CP-LOOM-jackknife+-APS", "LOOM, jackknife+, APS"), ("CP-PreVal-jackknife+", "PreVal, jackknife+, LAC"),
                   ("CP-PreVal-jackknife+-APS", "PreVal, jackknife+, APS"), ("CP-Split-TS", "Split conformal (TS), LAC"),
                   ("CP-Split-TS-APS", "Split conformal (TS), APS")]:
        try:
            rows.append({"Method": lab,
                         "Coverage, clean (mean / 10th pct.)": f"{f(cp(m, False))} / {f(cp(m, False, 'cov_q10'))}",
                         "Set size, clean": f2(cp(m, False, "size_mean")),
                         "Empty sets, clean": pct(cp(m, False, "empty_mean")),
                         "Coverage, corrupted": f(cp(m, True)),
                         "Set size, corrupted": f2(cp(m, True, "size_mean")),
                         "Empty sets, corrupted": pct(cp(m, True, "empty_mean"))})
        except KeyError:
            pass
    return md_table(pd.DataFrame(rows), caption)


def tab_eff(caption, methods):
    rows = []
    for m in methods:
        try:
            rows.append({"Method": NAMES.get(m, m), "Median fit time (s)": f(eff(m), 3), "Mean fit time (s)": f(eff(m, "fit_mean"), 2),
                         "Total fit time, all datasets (h)": f(eff(m, "fit_total") / 3600, 2),
                         "Median prediction time, 19 test sets (s)": "–" if pd.isna(eff(m, "pred_median")) else f(eff(m, "pred_median"), 2)})
        except KeyError:
            pass
    return md_table(pd.DataFrame(rows), caption)


def _allbetter_shift():
    """True only if LOOM is significantly better than every competitor in the corrupted-data NLL tables."""
    try:
        recs = S["E4_shifted"]["nll"]["pairwise"] + [r for r in S["E2b_raw_shifted"]["nll"]["pairwise"] if r["competitor"] != "LOOM-Ens5"]
        return all(r["p_holm"] < 0.05 and r["median_diff"] > 0 for r in recs)
    except KeyError:
        return False


ALLBETTER_SHIFT = _allbetter_shift()
CTX = {k: v for k, v in globals().items() if not k.startswith("__")}


NARRATIVE = ["ERROR_ANALYSIS", "DISCUSSION", "LIMITATIONS", "CONCLUSION", "ABSTRACT", "COVER_RESULT"]


def load_narrative():
    """Narrative paragraphs live in narr_<name>.md and may themselves contain {{ }} expressions."""
    for nm in NARRATIVE:
        p = HERE / f"narr_{nm.lower()}.md"
        CTX[nm] = render(p.read_text().strip()) if p.exists() else f"[{nm} — to be written after the final results]"


def render(text):
    def repl(mo):
        expr = mo.group(1).strip()
        try:
            return str(eval(expr, CTX))
        except Exception as e:
            if DRAFT:                      # draft builds mark missing values; final builds abort
                return f"**[MISSING: {expr[:60]}]**"
            raise RuntimeError(f"manuscript expression failed: {{{{{expr}}}}}: {e!r}")
    return re.sub(r"\{\{(.+?)\}\}", repl, text, flags=re.S)


def build(anonymous=True):
    load_narrative()
    parts = []
    for name in ["front.md", "sections_static.md", "sections_results.md", "REFS", "back.md"]:
        if name == "REFS":   # Springer: reference list first, then Statements and Declarations
            parts.append("# References {-}\n\n::: {#refs}\n:::\n")
            continue
        p = HERE / name
        if p.exists():
            parts.append(p.read_text())
    md = render("\n\n".join(parts))
    md = md.replace("](figures/", f"]({RES / 'figures'}/")
    (HERE / "manuscript_rendered.md").write_text(md)
    docx = OUT / ("Final_Manuscript.docx" if JOURNAL == "jksu" else "alternative_ESWA" / Path("Manuscript_ESWA.docx"))
    docx.parent.mkdir(exist_ok=True, parents=True)
    if JOURNAL == "eswa":   # Elsevier: highlights after the keywords, Harvard references
        hl = (HERE / "narr_highlights.md").read_text()
        md = md.replace("**Keywords**", "**Highlights**\n\n" + hl + "\n\n**Keywords**", 1)
        (HERE / "manuscript_rendered.md").write_text(md)
    cmd = ["pandoc", str(HERE / "manuscript_rendered.md"), "-o", str(docx), "--citeproc",
           f"--bibliography={HERE / 'references.json'}", f"--csl={HERE / ('springer-basic-author-date.csl' if JOURNAL == 'jksu' else 'elsevier-harvard.csl')}",
           f"--reference-doc={HERE / 'reference_springer.docx'}", "--resource-path", str(RES),
           "-M", "link-citations=false"]
    subprocess.run(cmd, check=True)
    postprocess(docx)
    print("written", docx)
    return docx


def postprocess(path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Cm
    d = Document(path)
    # page numbers in the footer
    for s in d.sections:
        p = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        for t, txt in (("begin", None), (None, " PAGE "), ("end", None)):
            if t:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); r._r.append(e)
            else:
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)
    # three-line tables, 8 pt, compact
    for tb in d.tables:
        tbl = tb._tbl
        pr = tbl.tblPr
        b = pr.find(qn("w:tblBorders"))
        if b is not None:
            pr.remove(b)
        b = OxmlElement("w:tblBorders")
        for edge in ("top", "bottom"):
            e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), "000000"); b.append(e)
        for edge in ("left", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "nil"); b.append(e)
        pr.append(b)
        for i, row in enumerate(tb.rows):
            for cell in row.cells:
                if i == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    bd = OxmlElement("w:tcBorders"); e = OxmlElement("w:bottom")
                    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "000000"); bd.append(e); tcPr.append(bd)
                for par in cell.paragraphs:
                    par.paragraph_format.space_after = Pt(0); par.paragraph_format.line_spacing = 1.0
                    for run in par.runs:
                        run.font.size = Pt(8); run.font.name = "Times New Roman"
                        if i == 0:
                            run.font.bold = True
    # figure paragraphs centred
    for par in d.paragraphs:
        if par._p.xpath(".//w:drawing"):
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.save(path)


def to_pdf(docx):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return docx.with_suffix(".pdf")


if __name__ == "__main__":
    dx = build()
    print("pdf:", to_pdf(dx))
